"""Reports that save themselves, as files the owner can open without the app.

Asked for by the owner: "once the system creates a report, make it autosave
in files". After every night audit the app writes, for that hotel:

* **`<date> Daily summary.pdf`** — the night's key figures (rooms, occupancy,
  ADR, RevPAR, revenue) for the day, month and year, the revenue by
  department, and the taxes and payments recorded;
* **`<YYYY-MM> Accountant pack.xlsx`** — the month so far: sales by line,
  taxes with their room-revenue base, and the guest, city and deposit ledgers.
  Rewritten each night, so it is always the month to date.

They go in `Documents › Open Hospitality › Saved reports › <code> – <hotel> ›
<YYYY-MM>`: one folder per hotel, one per month, named so they sort.

**Why a watcher and not a hook in ingestion.** A report reaches the books by
the drop folder, the upload page, the night-audit upload and (later) email.
Every one of those records the day it covered in `ingestion_coverage`, so
following that table is the one place that sees them all, and ingestion — the
engine — is not touched. A watermark in `desktop.setting` means a restart
writes what it missed and nothing twice.

**Nothing here writes to the books.** It reads the same statement the
dashboard shows. A file the owner has open in Excel cannot be replaced; that
night's file is retried on the next pass rather than lost.
"""

import logging
import re
import threading
from collections.abc import Callable, Iterable
from dataclasses import dataclass
from datetime import date, datetime
from decimal import Decimal
from pathlib import Path
from xml.sax.saxutils import escape

from sqlalchemy import func, select
from sqlalchemy.orm import Session

from usali import reporting
from usali.desktop.settings import read_setting, write_setting
from usali.models import IngestionCoverage, Property

_LOG = logging.getLogger(__name__)

WATERMARK_KEY = "saved_reports_coverage_id"
PROFILES_KEY = "hotel_profiles"  # welcome_api.PROFILES_KEY; not imported to keep this FastAPI-free

_UNSAFE = re.compile(r'[<>:"/\\|?*\x00-\x1f]')

_METRICS = (
    ("ROOMS_OCCUPIED", "Rooms occupied", "count"),
    ("TOTAL_ROOMS", "Rooms in the hotel", "count"),
    ("OCCUPANCY_PCT", "Occupancy", "pct"),
    ("ADR", "Average daily rate (ADR)", "money"),
    ("REVPAR", "Revenue per available room (RevPAR)", "money"),
    ("ROOM_REVENUE", "Room revenue", "money"),
)


def title_case(name: str) -> str:
    """The registry keeps names as reports print them — upper case. The same
    rule as the portal's `lib/propertyName.ts`."""
    return re.sub(r"(^|[\s(/-])([a-z])", lambda m: m.group(1) + m.group(2).upper(), name.lower())


@dataclass(frozen=True)
class Hotel:
    property_id: str
    name: str
    ownership_entity: str | None

    @property
    def folder_name(self) -> str:
        return _UNSAFE.sub("", f"{self.property_id} – {self.name}").strip(" .")


def hotel_folder(root: Path, hotel: Hotel) -> Path:
    return root / hotel.folder_name


def daily_summary_path(root: Path, hotel: Hotel, day: date) -> Path:
    return hotel_folder(root, hotel) / f"{day:%Y-%m}" / f"{day.isoformat()} Daily summary.pdf"


def accountant_pack_path(root: Path, hotel: Hotel, day: date) -> Path:
    return hotel_folder(root, hotel) / f"{day:%Y-%m}" / f"{day:%Y-%m} Accountant pack.xlsx"


# --- reading ------------------------------------------------------------------

def hotels(session: Session) -> dict[str, Hotel]:
    raw = read_setting(session, PROFILES_KEY)
    profiles = raw if isinstance(raw, dict) else {}
    return {
        p.property_id: Hotel(
            property_id=p.property_id,
            name=title_case(p.name),
            ownership_entity=(profiles.get(p.property_id) or {}).get("ownership_entity"),
        )
        for p in session.scalars(select(Property))
    }


def statement(session: Session, property_id: str, **when: date) -> reporting.SosReport:
    """The journal's statement, as the dashboard shows it; the facts' when the
    ledger is not posting (no fiscal calendar yet)."""
    try:
        return reporting.summary_operating_statement_from_journal(
            session, property_id=property_id, **when
        )
    except reporting.NoPostedEntriesError:
        return reporting.summary_operating_statement(session, property_id=property_id, **when)


# --- the daily summary (PDF) ----------------------------------------------------

def _money(v: Decimal | None) -> str:
    if v is None:
        return "—"
    return f"-${-v:,.2f}" if v < 0 else f"${v:,.2f}"


def _fmt(v: Decimal | None, kind: str) -> str:
    if v is None:
        return "—"
    if kind == "pct":
        return f"{v:.1f}%"
    if kind == "count":
        return f"{v:,.0f}"
    return _money(v)


def write_daily_summary(session: Session, hotel: Hotel, day: date, path: Path) -> None:
    from reportlab.lib import colors  # type: ignore[import-untyped]
    from reportlab.lib.pagesizes import letter  # type: ignore[import-untyped]
    from reportlab.lib.styles import getSampleStyleSheet  # type: ignore[import-untyped]
    from reportlab.lib.units import inch  # type: ignore[import-untyped]
    from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle  # type: ignore[import-untyped]

    report = statement(session, hotel.property_id, business_date=day)
    try:
        month = statement(session, hotel.property_id, date_from=day.replace(day=1), date_to=day)
        month_revenue: Decimal | None = month.total_operating_revenue
    except ValueError:
        month_revenue = None
    stats = {m.metric_code: m for m in report.statistics}

    styles = getSampleStyleSheet()
    grid = TableStyle([
        ("FONTNAME", (0, 0), (-1, 0), "Helvetica-Bold"),
        ("LINEBELOW", (0, 0), (-1, 0), 0.75, colors.HexColor("#16181D")),
        ("ALIGN", (1, 0), (-1, -1), "RIGHT"),
        ("FONTSIZE", (0, 0), (-1, -1), 9),
        ("BOTTOMPADDING", (0, 0), (-1, -1), 3),
        ("TOPPADDING", (0, 0), (-1, -1), 3),
        ("ROWBACKGROUNDS", (0, 1), (-1, -1), [colors.white, colors.HexColor("#F4F4F2")]),
    ])
    story: list[object] = [
        # Paragraphs are markup: "Holiday Inn & Suites" must be escaped.
        Paragraph(f"{escape(hotel.name)} <font color='#6B6F76'>· {escape(hotel.property_id)}</font>",
                  styles["Title"]),
        Paragraph(
            f"Daily summary for <b>{day:%A, %B} {day.day}, {day.year}</b>"
            + (f" &nbsp;·&nbsp; {escape(hotel.ownership_entity)}" if hotel.ownership_entity else ""),
            styles["Normal"],
        ),
        Spacer(1, 0.2 * inch),
        Paragraph("Key figures", styles["Heading2"]),
    ]
    rows = [["", "Day", "Month to date", "Year to date"]]
    for code, label, kind in _METRICS:
        m = stats.get(code)
        if m is not None:
            rows.append([label, _fmt(m.day, kind), _fmt(m.mtd, kind), _fmt(m.ytd, kind)])
    rows.append(["Total operating revenue", _money(report.total_operating_revenue),
                 _money(month_revenue), "—"])
    story.append(Table(rows, colWidths=[2.8 * inch, 1.3 * inch, 1.3 * inch, 1.3 * inch],
                       style=grid, hAlign="LEFT"))

    def section(title: str, lines: Iterable[tuple[str, Decimal]], total: Decimal) -> None:
        body = [[name, _money(amount)] for name, amount in lines]
        if not body:
            return
        story.extend([Spacer(1, 0.15 * inch), Paragraph(title, styles["Heading2"])])
        body = [["", "Day"], *body, ["Total", _money(total)]]
        table = Table(body, colWidths=[4.6 * inch, 2.1 * inch], style=grid, hAlign="LEFT")
        table.setStyle(TableStyle([("FONTNAME", (0, -1), (-1, -1), "Helvetica-Bold"),
                                   ("LINEABOVE", (0, -1), (-1, -1), 0.5, colors.grey)]))
        story.append(table)

    section(
        "Revenue by department",
        [(f"{d.sub_category} — {line.line_item}", line.total)
         for d in report.operated_departments for line in d.lines]
        + [(f"Miscellaneous — {line.line_item}", line.total) for line in report.misc_income],
        report.total_operating_revenue,
    )
    section("Taxes collected", [(line.line_item, line.total) for line in report.taxes],
            report.taxes_total)
    section("Payments and settlements",
            [(line.line_item, line.total) for line in report.settlements],
            report.settlements_total)

    story.extend([
        Spacer(1, 0.25 * inch),
        Paragraph(
            f"<font size='8' color='#6B6F76'>From the night audit reports read by Open "
            f"Hospitality, as the books stood at {datetime.now():%Y-%m-%d %H:%M}. Not an "
            f"audited statement.</font>",
            styles["Normal"],
        ),
    ])
    _write_atomically(path, lambda tmp: SimpleDocTemplate(
        str(tmp), pagesize=letter, title=f"{hotel.name} daily summary {day.isoformat()}",
        author="Open Hospitality", leftMargin=0.75 * inch, rightMargin=0.75 * inch,
        topMargin=0.7 * inch, bottomMargin=0.7 * inch,
    ).build(story))


# --- the accountant pack (Excel) ---------------------------------------------

def write_accountant_pack(session: Session, hotel: Hotel, day: date, path: Path) -> None:
    from openpyxl import Workbook  # type: ignore[import-untyped]
    from openpyxl.styles import Font  # type: ignore[import-untyped]

    pack = reporting.cpa_pack(session, property_id=hotel.property_id, month=f"{day:%Y-%m}")
    book = Workbook()
    bold = Font(bold=True)
    money = "#,##0.00;[Red]-#,##0.00"

    def sheet(title: str, header: list[str], rows: list[list[object]], first: bool = False) -> None:
        ws = book.active if first else book.create_sheet()
        assert ws is not None
        ws.title = title
        ws.append(header)
        for cell in ws[1]:
            cell.font = bold
        for row in rows:
            ws.append(row)
        for column in ws.columns:
            width = max(len(str(c.value or "")) for c in column)
            ws.column_dimensions[column[0].column_letter].width = min(max(width + 2, 10), 60)
            for c in column[1:]:
                if isinstance(c.value, Decimal | float | int) and not isinstance(c.value, bool):
                    c.number_format = money

    sheet("Summary", ["Accountant pack", "Open Hospitality"], [
        ["Hotel", hotel.name],
        ["Hotel code", hotel.property_id],
        ["Ownership entity", hotel.ownership_entity or ""],
        ["Month", pack.month],
        ["Through", day.isoformat()],
        ["Total operating revenue", pack.sales.total_operating_revenue],
        ["Taxes collected", pack.taxes.taxes_total],
        ["Room revenue (occupancy-tax base)", pack.taxes.room_revenue_base],
    ], first=True)
    sheet("Sales", ["Category", "Department", "Line", "Month to date", "Days"], [
        [s.major, s.sub_category, s.line_item, s.mtd_amount, s.day_count]
        for s in pack.sales.lines
    ] + [["Total operating revenue", "", "", pack.sales.total_operating_revenue, ""]])
    sheet("Taxes", ["Tax", "GL account", "Month to date"], [
        [t.line_item, t.gl_account_code or "", t.mtd_amount] for t in pack.taxes.lines
    ] + [["Taxes total", "", pack.taxes.taxes_total],
         ["Room revenue base", "", pack.taxes.room_revenue_base]])
    # "First reported", not "opening": the earliest balance IN the month, which
    # is not the prior month's close after a gap (reporting.ArLine).
    sheet("Ledgers", ["Ledger", "Code", "First reported", "Latest", "Movement"], [
        [a.ledger_name, a.ledger_code, a.opening_balance, a.closing_balance, a.movement]
        for a in pack.ar.balances
    ])
    _write_atomically(path, lambda tmp: book.save(str(tmp)))


def _write_atomically(path: Path, write: Callable[[Path], object]) -> None:
    """Write beside, then replace: an owner opening the file mid-write sees
    the old one, never half a file. A file they have OPEN cannot be replaced
    on Windows — that raises, and the caller retries next pass."""
    path.parent.mkdir(parents=True, exist_ok=True)
    tmp = path.with_name(f"~{path.stem}.writing{path.suffix}")
    try:
        write(tmp)
        tmp.replace(path)
    finally:
        tmp.unlink(missing_ok=True)


# --- the watcher --------------------------------------------------------------

def save_new(session_factory: Callable[[], Session], root: Path) -> list[Path]:
    """Write the files for every day read since the last pass. Returns what
    was written. The watermark only moves past days whose files were written."""
    written: list[Path] = []
    with session_factory() as session:
        mark = read_setting(session, WATERMARK_KEY)
        after = mark if isinstance(mark, int) else 0
        new = session.execute(
            select(IngestionCoverage.property_id, IngestionCoverage.business_date,
                   func.max(IngestionCoverage.coverage_id))
            .where(IngestionCoverage.coverage_id > after)
            .group_by(IngestionCoverage.property_id, IngestionCoverage.business_date)
        ).all()
        if not new:
            return written
        known = hotels(session)
        failed_at: int | None = None
        for property_id, day, coverage_id in sorted(new, key=lambda r: r[2]):
            hotel = known.get(property_id)
            if hotel is None:
                continue
            try:
                for target, write in (
                    (daily_summary_path(root, hotel, day), write_daily_summary),
                    (accountant_pack_path(root, hotel, day), write_accountant_pack),
                ):
                    write(session, hotel, day, target)
                    written.append(target)
            except reporting.NoFactsError:
                continue  # a report with nothing to summarise (statistics-only days, say)
            except (OSError, ValueError) as exc:
                _LOG.warning("could not save %s %s yet: %s", property_id, day, exc)
                failed_at = coverage_id if failed_at is None else min(failed_at, coverage_id)
            finally:
                session.rollback()  # reads only; never leave a transaction open
        top = max(r[2] for r in new)
        advance = top if failed_at is None else failed_at - 1
        if advance > after:
            write_setting(session, WATERMARK_KEY, advance)
            session.commit()
    return written


class SavedReports:
    """Runs `save_new` every `interval` seconds in the background."""

    def __init__(self, session_factory: Callable[[], Session], root: Path,
                 interval: float = 30.0) -> None:
        self.session_factory = session_factory
        self.root = root
        self.interval = interval
        self._stop = threading.Event()
        self._thread: threading.Thread | None = None

    def start(self) -> None:
        self.root.mkdir(parents=True, exist_ok=True)
        self._thread = threading.Thread(target=self._run, name="saved-reports", daemon=True)
        self._thread.start()

    def _run(self) -> None:
        while True:
            try:
                for path in save_new(self.session_factory, self.root):
                    _LOG.info("saved %s", path.name)
            except Exception:
                _LOG.exception("saving reports failed; will try again")
            if self._stop.wait(self.interval):
                return

    def stop(self) -> None:
        self._stop.set()
        if self._thread is not None:
            self._thread.join(timeout=10)
